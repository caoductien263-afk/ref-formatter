"""
Gemini service module for Streamlit Gemini Reference Formatter.
Encapsulates API key resolution, multimodal file uploads (PDF File API, Word via python-docx),
and prompt orchestration for guideline analysis and reference formatting.
"""

import os
import time
import tempfile
import io
from typing import Optional, Union, List, Any


def resolve_api_key(user_input_key: Optional[str] = None) -> Optional[str]:
    """
    Resolves the Gemini API key strictly from the User UI input.
    Ignores secrets to ensure no old invalid keys are picked up.
    """
    if user_input_key and user_input_key.strip():
        return user_input_key.strip()
        
    return None


def extract_docx_text(file_bytes: bytes) -> str:
    """
    Extracts text and table content from a .docx binary file using python-docx.
    Avoids Gemini API 400 Unsupported MIME type error for Word files.
    """
    import docx
    
    doc = docx.Document(io.BytesIO(file_bytes))
    extracted_paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    
    # Extract table rows if present (academic guidelines often contain formatting tables)
    table_lines = []
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                # Deduplicate repeated text from merged cells
                deduped = []
                for cell_val in row_cells:
                    if not deduped or deduped[-1] != cell_val:
                        deduped.append(cell_val)
                table_lines.append(" | ".join(deduped))
                
    full_text_parts = []
    if extracted_paragraphs:
        full_text_parts.append("\n".join(extracted_paragraphs))
    if table_lines:
        full_text_parts.append("[Bảng biểu trích xuất từ tài liệu]:\n" + "\n".join(table_lines))
        
    return "\n\n".join(full_text_parts)


def extract_guideline_from_file_or_text(
    api_key: str,
    text: Optional[str] = None,
    file_bytes: Optional[bytes] = None,
    filename: Optional[str] = None,
    model_name: str = "gemini-1.5-pro",
) -> str:
    """
    Extracts and synthesizes academic reference rules across 8 standard dimensions.
    Handles PDF files via Gemini File API with state polling and guaranteed cleanup.
    Handles Word (.docx) files via local text extraction to prevent 400 MIME errors.
    """
    if not api_key or not api_key.strip():
        raise ValueError("Chưa có API Key. Vui lòng cung cấp Google Gemini API Key.")
        
    has_text = bool(text and text.strip())
    has_file = bool(file_bytes and len(file_bytes) > 0)
    
    if not has_text and not has_file:
        raise ValueError("Vui lòng cung cấp văn bản quy chuẩn hoặc tải lên tệp tài liệu mẫu (PDF/Word).")

    import google.generativeai as genai
    genai.configure(api_key=api_key.strip())

    prompt = """Bạn là một Chuyên gia Biên tập & Chuẩn hóa Tài liệu Tham khảo Khoa học.
Nhiệm vụ của bạn là phân tích tài liệu/văn bản quy chuẩn được cung cấp và rút ra bộ quy tắc định dạng bắt buộc cho phần Tài liệu tham khảo (References).

Hãy phân tích kỹ lưỡng và trình bày tóm tắt rõ ràng theo đúng 8 chiều kích thước chuẩn sau:
1. Tạp chí khoa học (Journal Articles): Cấu trúc, vị trí tên tác giả, năm, tên bài báo, tên tạp chí (in nghiêng?), tập/số (volume/issue), số trang (pages), mã DOI.
2. Sách và chương sách (Books & Book Chapters): Cấu trúc, tác giả/chủ biên, năm xuất bản, tên sách (in nghiêng?), nhà xuất bản, nơi xuất bản, số trang.
3. Kỷ yếu hội thảo/hội nghị (Conference Proceedings): Cấu trúc, tên tác giả, năm, tên bài viết, tên hội nghị, địa điểm tổ chức, nhà xuất bản.
4. Website và tài liệu trực tuyến (Websites & Reports): Cấu trúc, tác giả/tổ chức, tên bài viết, đường dẫn URL, ngày truy cập (nếu có yêu cầu).
5. Quy tắc tên tác giả (Authors format): Thứ tự Họ - Tên hay Tên - Họ; cách viết tắt chữ lót/tên; quy tắc khi có 1 tác giả, 2-3 tác giả, và từ 4 tác giả trở lên (có dùng "et al." không, dùng sau bao nhiêu tác giả).
6. Quy tắc năm xuất bản (Publication Year): Vị trí đặt năm (ngay sau tên tác giả hay cuối mục); có đặt trong dấu ngoặc đơn () hay không; có dấu chấm phía sau không.
7. Quy tắc tiêu đề (Titles formatting): Tiêu đề bài báo/sách/tạp chí viết hoa chữ cái đầu (Sentence case) hay viết hoa từng từ (Title case); có đặt trong ngoặc kép "" hay in nghiêng *...*.
8. Quy tắc dấu câu và các trường thông tin khác: Dấu chấm, phẩy, hai chấm giữa các trường; cách viết số tập (vol), số phát hành (no/issue), khoảng trang (pp. hay chỉ số trang).

Đưa ra ví dụ minh họa chuẩn xác cho từng loại tài liệu dựa trên đúng các quy tắc vừa tổng hợp."""

    contents: List[Any] = [prompt]
    tmp_path: Optional[str] = None
    gemini_file: Optional[Any] = None

    try:
        # Process attached file if present
        if has_file and filename:
            ext = os.path.splitext(filename)[1].lower()
            
            if ext == ".docx":
                # Local Word extraction via python-docx
                docx_content = extract_docx_text(file_bytes)  # type: ignore[arg-type]
                contents.append(f"\n\n--- NỘI DUNG TỪ TỆP WORD ({filename}) ---\n{docx_content}")
            elif ext == ".pdf":
                # PDF upload via Gemini File API
                with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
                    tmp.write(file_bytes)  # type: ignore[arg-type]
                    tmp_path = tmp.name
                
                # Upload to Gemini File API
                gemini_file = genai.upload_file(path=tmp_path, mime_type="application/pdf")
                
                # Poll until file is ACTIVE
                poll_start = time.time()
                while gemini_file.state.name == "PROCESSING":
                    time.sleep(1)
                    gemini_file = genai.get_file(gemini_file.name)
                    if time.time() - poll_start > 60:
                        raise TimeoutError("Quá thời gian chờ xử lý file PDF trên Gemini API (Timeout 60s).")
                
                if gemini_file.state.name == "FAILED":
                    raise RuntimeError(f"Xử lý file PDF trên Gemini thất bại (State: {gemini_file.state.name}).")
                
                contents.append(gemini_file)
            else:
                # Text/plain or other format fallback
                try:
                    plain_text = file_bytes.decode("utf-8")  # type: ignore[union-attr]
                    contents.append(f"\n\n--- NỘI DUNG TỆP ({filename}) ---\n{plain_text}")
                except Exception:
                    raise ValueError(f"Định dạng tệp '{ext}' không được hỗ trợ. Vui lòng tải file PDF hoặc DOCX.")

        # Process user guideline text if provided
        if has_text:
            contents.append(f"\n\n--- VĂN BẢN QUY CHUẨN ĐƯỢC CUNG CẤP ---\n{text}")

        model = genai.GenerativeModel(model_name)
        response = model.generate_content(contents)
        return response.text

    finally:
        # Deterministic cleanup of ephemeral files (both local disk and remote cloud)
        if gemini_file:
            try:
                genai.delete_file(gemini_file.name)
            except Exception:
                pass
        if tmp_path and os.path.exists(tmp_path):
            try:
                os.remove(tmp_path)
            except Exception:
                pass


def format_references(
    api_key: str,
    guideline_summary: str,
    raw_references: str,
    model_name: str = "gemini-1.5-pro",
) -> str:
    """
    Formats raw references according to approved guideline rules.
    Enforces double newlines (\\n\\n) between entries, Markdown italics (*...*),
    order preservation, and explicit anti-hallucination tag [Điền ngày/tháng/năm].
    """
    if not api_key or not api_key.strip():
        raise ValueError("Chưa có API Key. Vui lòng cung cấp Google Gemini API Key.")
    if not guideline_summary or not guideline_summary.strip():
        raise ValueError("Quy chuẩn tạp chí trống. Vui lòng hoàn thành Bước 1 và 2 trước.")
    if not raw_references or not raw_references.strip():
        raise ValueError("Danh sách tài liệu tham khảo thô trống. Vui lòng nhập dữ liệu.")

    import google.generativeai as genai
    genai.configure(api_key=api_key.strip())

    prompt = f"""Bạn là một Chuyên gia Biên tập & Chuẩn hóa Tài liệu Tham khảo Khoa học.
Dựa vào bộ quy chuẩn đã được phê duyệt sau đây:

{guideline_summary}

Hãy chuyển đổi toàn bộ danh sách Tài liệu Tham khảo (TLTK) thô dưới đây sang đúng định dạng chuẩn trên.

CÁC NGUYÊN TẮC BẮT BUỘC:
1. GIỮ NGUYÊN HOÀN TOÀN THỨ TỰ ban đầu của các tài liệu (từ mục đầu tiên đến hết theo đúng danh sách thô). Tuyệt đối không tự ý đảo thứ tự theo bảng chữ cái trừ khi bộ quy chuẩn yêu cầu rõ ràng.
2. ĐẢM BẢO ĐỘ CHUẨN XÁC CAO: Chuẩn từng dấu chấm, dấu phẩy, khoảng trắng, viết hoa/viết thường theo quy định.
3. ĐỊNH DẠNG IN NGHIÊNG: Dùng cú pháp Markdown *tên in nghiêng* (ví dụ: *Journal of Science*, *Tên Sách*) để người dùng có thể bôi đen và sao chép trực tiếp vào Microsoft Word mà vẫn giữ nguyên in nghiêng.
4. MỖI TÀI LIỆU LÀ MỘT ĐOẠN ĐỘC LẬP: Phải ngăn cách giữa các tài liệu bằng đúng hai ký tự xuống dòng (\\n\\n). Tuyệt đối không dồn nhiều tài liệu vào một đoạn văn.
5. NGUYÊN TẮC CHỐNG TỰ BỊA ĐẶT THÔNG TIN (ANTI-HALLUCINATION):
   - TUYỆT ĐỐI KHÔNG tự bịa đặt năm xuất bản, ngày tháng truy cập, số trang, số tập, nhà xuất bản, mã DOI nếu tài liệu thô không có.
   - Nếu tài liệu thiếu ngày/tháng/năm xuất bản hoặc ngày truy cập website, BẮT BUỘC đánh dấu bằng đúng cụm từ: [Điền ngày/tháng/năm].
   - Nếu thiếu các thông tin khác mà quy chuẩn yêu cầu (như số trang, DOI), hãy đánh dấu rõ: [Điền số trang], [Điền DOI].
6. ĐỊNH DẠNG ĐẦU MỤC: Nếu danh sách thô có đánh số [1], [2]... hoặc quy chuẩn yêu cầu đánh số, hãy đánh số liên tục bắt đầu từ [1].
7. KHÔNG XUẤT CODE BLOCK: Chỉ trả về danh sách tài liệu tham khảo hoàn chỉnh, không bao bọc kết quả bằng khối mã markdown (```markdown hoặc ```). Không thêm lời chào, mở bài hoặc kết luận.

DANH SÁCH TLTK THÔ CẦN ĐỊNH DẠNG:
{raw_references}
"""

    model = genai.GenerativeModel(model_name)
    response = model.generate_content(prompt)
    return response.text
